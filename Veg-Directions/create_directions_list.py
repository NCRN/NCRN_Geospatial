#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Initial working script for creating directions to veg monitoring locations.
--------------------------------------------------------------------------------
TODO: Add more of a complete description.

References:
# https://betterprogramming.pub/automate-filling-templates-with-python-1ff6c6fd595e
"""

#__author__ = "NCRN GIS Unit"
#__copyright__ = "None"
#__credits__ = [""]
#__license__ = "GPL"
#__version__ = "1.0.2"
#__maintainer__ = "David Jones"
#__email__ = "david_jones@nps.gov"
#__status__ = "Staging"

# Import statements for utilized libraries / packages
from cmath import nan
from tkinter.tix import ROW
import openpyxl
from openpyxl import load_workbook
import docx
from docx import Document
import docxtpl
from docxtpl import DocxTemplate
import os
import pandas as pd

print("### GETTING STARTED ###".format())

"""
Set various global variables. Some of these could be parameterized to be used in an 
ArcGIS Toolbox script and/or command line use. 
"""
# Currently hardcoded values that may be parameterized if bundling into a tool

main_path = r'C:\Users\goettel\DOI\NPS-NCRN-Forest Veg - Documents\General\Field_Maps_2023'
template_path = r'C:\Users\goettel\DOI\NCRN Data Management - Geospatial\Templates\NCRN\Directions_Template.docx' ## Create a variable to store the full path to the Word doc template. Set prefix to the user portion of the root directory
__XCEL_LIBRARY = r'C:\Users\goettel\DOI\NCRN Data Management - Geospatial\NCRN_Veg_Locations_Directions_Master.xlsx' ## Create a variable to store the full path to the Excel file. Set prefix to the user portion of the root directory

workbook = load_workbook(__XCEL_LIBRARY, data_only=True)
template = DocxTemplate(template_path)

df_areas = pd.read_excel(__XCEL_LIBRARY, sheet_name='Areas Python')
df_plots = pd.read_excel(__XCEL_LIBRARY, sheet_name='Plots', na_filter = False)

for index, row in df_areas.iterrows():
    Area_Name = row['Area Name']
    if row['Area Warnings'] == 0:
        Area_Warnings = ""
    else: 
        Area_Warnings = row['Area Warnings']
    to_fill_in_area = {
        'Area_Name': Area_Name,
        'Area_Directions': row['Area Directions'],
        'Area_Warnings': Area_Warnings,
        'Plot_Name': '{{Plot_Name}}',
        'Plot_Directions': '{{Plot_Directions}}',
        'Warnings': '{{Warnings}}',
        'Keys': '{{Keys}}',
        'Map_N': '{{Map_N}}'
        }
    template.render(to_fill_in_area)
    filename = str(to_fill_in_area['Area_Name']) + ' Directions.docx'
    filled_path = os.path.join(main_path, row['Panel'], filename)
    template.save(filled_path)
    print("Done with creating %s" % str(to_fill_in_area['Area_Name']))
    for index, row in df_plots.iterrows():
        filled_template = DocxTemplate(filled_path)
        if row['Area/Route - Panel'] == Area_Name:
            to_fill_in_plot = {
                'Plot_Name': row['Plot Name'],
                'Plot_Directions': row['Plot Directions'],
                'Warnings': row['Plot Warnings'],
                'Keys': row['Keys'],
                'Map_N': row['Map #']
                }
            filled_template.render(to_fill_in_plot)
            table = filled_template.tables[0]
            row = table.add_row().cells
            row[0].text = '{{Plot_Name}}'
            row[1].text = '{{Plot_Directions}}'
            row[2].text = '{{Warnings}}'
            row[3].text = '{{Keys}}'
            row[4].text = '{{Map_N}}'
            filled_template.save(filled_path)
        to_fill_in = {}
    filled_template.render(to_fill_in)
    filled_template.save(filled_path)
    print("Done adding plots to %s" % str(to_fill_in_area['Area_Name']))